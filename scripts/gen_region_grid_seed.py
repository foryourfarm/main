"""기상청 배포 격자 엑셀 → docs/seed/region_grid_seed.csv 생성.

기상청 단기예보 API는 격자좌표(nx, ny)만 받는다. 그 대응은 기상청이 배포하는
"단기예보 조회서비스 오픈API활용가이드 격자_위경도" 엑셀이 권위 있는 정답이므로,
위경도를 지오코딩해 계산하지 않고 이 표를 그대로 쓴다.

왜 계산하지 않는가: 위경도→격자 변환기(app/infra/public_api/kma_grid.py)는 이 엑셀의
읍면동 3,564건과 교차검증해 98.85% 일치했으나, 격자 경계에 걸친 41건에서 1칸씩 어긋난다.
엑셀이 있는 지역은 엑셀 값을 쓰는 것이 정확하다(변환기는 엑셀에 없는 좌표용 폴백).

또한 지오코딩 경로는 쓸 수 없다 — VWorld는 "별도의 저장장치나 데이터베이스에 저장할 수
없음"이 약관에 명시돼 있어 region_grid 적재 목적에 부적합하다.

엑셀은 2026-07-01 기준으로 행정통합 신규 코드를 이미 반영하고 있어(순천시 1215000000)
우리 region_seed의 bjd_code와 256/256 직접 매칭된다 — 통합 전 코드 변환이 필요 없다.

사용법(엑셀은 리포에 커밋하지 않는다 — 산출 CSV만 커밋):
    backend/.venv/Scripts/python.exe scripts/gen_region_grid_seed.py "<격자_위경도.xlsx>"
"""

import csv
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SEED_DIR = ROOT / "docs" / "seed"
REGION_SEED = SEED_DIR / "region_seed.csv"
OUT_CSV = SEED_DIR / "region_grid_seed.csv"

COL_CODE = "행정구역코드"
COL_SIDO = "1단계"
COL_SGG = "2단계"
COL_EMD = "3단계"
COL_NX = "격자 X"
COL_NY = "격자 Y"


def main() -> None:
    if len(sys.argv) != 2:
        print('사용법: python scripts/gen_region_grid_seed.py "<격자_위경도.xlsx>"')
        raise SystemExit(1)

    wb = openpyxl.load_workbook(Path(sys.argv[1]), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    idx = {name: i for i, name in enumerate(rows[0]) if name}
    # 시군구 레벨 = 2단계 있고 3단계 없는 행. 읍면동 행은 region 단위가 아니라 제외.
    grid_by_code: dict[str, tuple[int, int]] = {}
    for r in rows[1:]:
        if not r[idx[COL_SGG]] or r[idx[COL_EMD]]:
            continue
        try:
            grid_by_code[str(r[idx[COL_CODE]])] = (int(r[idx[COL_NX]]), int(r[idx[COL_NY]]))
        except (TypeError, ValueError):
            continue

    seed = list(REGION_SEED.open(encoding="utf-8"))
    reader = csv.DictReader(seed)
    out_rows: list[dict[str, str]] = []
    missing: list[str] = []
    for s in reader:
        grid = grid_by_code.get(s["bjd_code"])
        if grid is None:
            missing.append(f"{s['sido']} {s['name']}({s['bjd_code']})")
            continue
        out_rows.append(
            {
                "region_id": s["id"],
                "name": s["name"],
                "sido": s["sido"],
                "nx": str(grid[0]),
                "ny": str(grid[1]),
            }
        )

    with OUT_CSV.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["region_id", "name", "sido", "nx", "ny"])
        w.writeheader()
        w.writerows(out_rows)

    print(f"격자 시드 {len(out_rows)}건 -> {OUT_CSV.relative_to(ROOT)}")
    if missing:
        # 조용히 빠지면 그 지역 단기예보가 안 되므로 반드시 드러낸다.
        print(f"[확인 필요] 엑셀에 없는 지역 {len(missing)}건: {missing}")


if __name__ == "__main__":
    main()
