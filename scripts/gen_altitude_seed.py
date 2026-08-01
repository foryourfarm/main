"""구역·법정동 대표 고도 시드 생성 (오프라인·재실행 가능, 앱 밖).

**왜 필요한가**: 기온은 고도에 직접 지배된다(환경감률 약 0.65℃/100m). 실측(10km 이내
지점쌍, 7월 일최고기온)에서 고도차 200~500m면 MAE 2.29℃, 500m 넘으면 5.07℃다. 거리를
좁혀도 안 잡힌다 — 평년치 도너를 거리만으로 고르면 한라산 지점과 해안 지점이 같은 후보에
든다(서귀포시 도너 고도 산포 1,524m). 고도를 알아야 도너를 거르고, 밭 위치로 감률 보정을
한다. nexttodo.md "구역·밭 고도 확보" 참고.

**표고 출처: SRTM 30m (NASA, 공공도메인)** — OpenTopoData 공개 API로 조회하고 **값만**
시드 CSV에 커밋한다(런타임 의존 없음). 대안 검토 결과:
  - VWorld: 고도 Open API가 2019년 종료됐고, 약관상 저장 금지라 시드 용도에 부적합
    (`gen_region_grid_seed.py`가 지오코딩을 버린 것과 같은 이유).
  - 국토지리정보원 DEM(data.go.kr 15059920, 이용허락 제한 없음): 5m급으로 가장 정확하나
    IMG 래스터 도엽 단위 + 로그인 다운로드 + GDAL 계열 의존성이 필요하다.
  - SRTM을 **우리 관측지점 752개의 실측 고도**에 대고 검증했다:
        MAE 8.8m / median 3.0m / p90 19m / bias +1.7m
    감률 환산 **0.06℃**다. 잡으려는 오차가 2.29~5.07℃라 DEM 5m급으로 올려 얻는
    이득(0.06→0.01℃)이 의미가 없다. 100m 넘게 튄 10개는 급경사지 지점 좌표 반올림으로 보인다.

**좌표 출처: 기상청 격자 엑셀** — 이미 리포 밖에 갖고 있는 파일이고 추가 의존성이 0이다.
  - 시군구(2단계) 행: 행정구역코드가 region_seed.bjd_code와 256/256 직매칭된다.
  - 읍면동(3단계) 행은 **행정동 코드**라 법정동 코드 직매칭이 1,408건뿐이다. 그래서
    (시도, 시군구, 읍면동) **이름**으로 붙인다 — 20,275건 중 17,123건(84.5%)이 붙고
    못 붙는 3,152건은 전부 도시 법정동(예: 종로구 청운동)이다. 읍·면·리는 다 붙는다.
  - 못 붙은 행은 시군구 대표점으로 폴백하고 `altitude_source`에 그 사실을 남긴다(§18-4).

**왜 격자중심이 아니라 엑셀 대표점인가**: 격자중심(5km 격자)은 산비탈에 떨어진다.
엑셀 대표점(관청 소재지)은 사람·농지가 있는 저지대다 — 우리가 원하는 "구역 대표 농지 고도"에
가깝다. 실측 비교(대표점 / 격자중심): 서귀포시 76m/209m, 정선군 309m/499m, 문경시 82m/197m.

**한계**: 리(里)는 소속 읍·면 대표점 고도를 쓴다. 산간 면은 면 내 기복이 300m를 넘어
(≈2℃) 이 근사가 남는 오차의 주 원인이다. 근본 해결은 법정구역 SHP의 리 단위 중심점이고,
그때 바뀌는 건 **이 스크립트뿐**이다 — 소비처는 좌표를 보지 않고 `altitude_m`만 본다.

실행:
    backend/.venv/Scripts/python.exe scripts/gen_altitude_seed.py "<격자_위경도.xlsx>"
"""

import csv
import json
import sys
import time
import urllib.request
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SEED_DIR = ROOT / "docs" / "seed"
REGION_SEED = SEED_DIR / "region_seed.csv"
BJD_SEED = SEED_DIR / "bjd_to_region.csv"
REGION_OUT = SEED_DIR / "region_altitude_seed.csv"
DISTRICT_OUT = SEED_DIR / "district_altitude_seed.csv"

# OpenTopoData 공개 인스턴스 제한: 100 locations/call, 1 call/sec, 1000 call/day.
# 좌표를 중복 제거하면 약 3,800점 → 38콜이라 제한 안쪽이다.
API = "https://api.opentopodata.org/v1/srtm30m"
BATCH = 100
SLEEP_SEC = 1.2

SOURCE_EMD = "emd_point"  # 읍·면·동 대표점 (리는 소속 읍·면 대표점)
SOURCE_REGION = "region_point"  # 시군구 대표점 폴백 (도시 법정동)

COL_CODE, COL_SIDO, COL_SGG, COL_EMD = 1, 2, 3, 4
COL_LON, COL_LAT = 13, 14


def read_excel(path: Path) -> tuple[dict[str, tuple[float, float]], dict[tuple, tuple[float, float]]]:
    """(행정구역코드→시군구 좌표, (시도,시군구,읍면동)→좌표)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
    wb.close()

    by_code: dict[str, tuple[float, float]] = {}
    by_name: dict[tuple, tuple[float, float]] = {}
    for r in rows[1:]:
        try:
            point = (float(r[COL_LAT]), float(r[COL_LON]))
        except (TypeError, ValueError):
            continue
        if r[COL_EMD]:
            by_name[(r[COL_SIDO], r[COL_SGG], r[COL_EMD])] = point
        elif r[COL_SGG]:
            by_code[str(r[COL_CODE])] = point
    return by_code, by_name


def fetch_altitudes(points: list[tuple[float, float]]) -> dict[tuple[float, float], float]:
    """좌표 → 표고(m). 좌표를 중복 제거해 부르므로 호출 수가 지점 수보다 훨씬 적다(§18-1)."""
    out: dict[tuple[float, float], float] = {}
    for i in range(0, len(points), BATCH):
        chunk = points[i : i + BATCH]
        query = "|".join(f"{lat},{lon}" for lat, lon in chunk)
        with urllib.request.urlopen(f"{API}?locations={query}", timeout=60) as resp:
            results = json.load(resp)["results"]
        for point, res in zip(chunk, results):
            # 해상 격자 등은 null이 온다. 값을 지어내지 않고 비운다(§18-4).
            if res.get("elevation") is not None:
                out[point] = float(res["elevation"])
        print(f"  표고 {min(i + BATCH, len(points))}/{len(points)}", flush=True)
        time.sleep(SLEEP_SEC)
    return out


def main() -> None:
    if len(sys.argv) != 2:
        print('사용법: python scripts/gen_altitude_seed.py "<격자_위경도.xlsx>"')
        raise SystemExit(1)

    by_code, by_name = read_excel(Path(sys.argv[1]))
    regions = list(csv.DictReader(REGION_SEED.open(encoding="utf-8")))
    districts = list(csv.DictReader(BJD_SEED.open(encoding="utf-8")))

    # 구역 좌표
    region_point: dict[str, tuple[float, float]] = {}
    for r in regions:
        point = by_code.get(r["bjd_code"])
        if point is not None:
            region_point[r["id"]] = point
    missing_regions = [r["name"] for r in regions if r["id"] not in region_point]

    # 법정동 좌표 — 이름 매칭, 실패 시 시군구 폴백
    region_id_by_name = {(r["sido"], r["name"]): r["id"] for r in regions}
    district_point: dict[str, tuple[tuple[float, float], str]] = {}
    for d in districts:
        point = by_name.get((d["sido"], d["region_name"], d["eupmyeondong"]))
        if point is not None:
            district_point[d["bjd_code"]] = (point, SOURCE_EMD)
            continue
        region_id = region_id_by_name.get((d["sido"], d["region_name"]))
        fallback = region_point.get(region_id) if region_id else None
        if fallback is not None:
            district_point[d["bjd_code"]] = (fallback, SOURCE_REGION)

    unique = sorted({*region_point.values(), *(p for p, _ in district_point.values())})
    print(f"구역 {len(region_point)}/{len(regions)}개, 법정동 {len(district_point)}/{len(districts)}개")
    print(f"중복 제거 좌표 {len(unique)}점 → API {(len(unique) + BATCH - 1) // BATCH}콜\n")

    altitude = fetch_altitudes(unique)

    with REGION_OUT.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["region_id", "altitude_m"])
        w.writeheader()
        for region_id, point in sorted(region_point.items(), key=lambda t: int(t[0])):
            if point in altitude:
                w.writerow({"region_id": region_id, "altitude_m": round(altitude[point])})

    by_source: dict[str, int] = {}
    with DISTRICT_OUT.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["bjd_code", "altitude_m", "altitude_source"])
        w.writeheader()
        for bjd_code, (point, source) in sorted(district_point.items()):
            if point not in altitude:
                continue
            by_source[source] = by_source.get(source, 0) + 1
            w.writerow(
                {
                    "bjd_code": bjd_code,
                    "altitude_m": round(altitude[point]),
                    "altitude_source": source,
                }
            )

    print(f"\n{REGION_OUT.relative_to(ROOT)}  {sum(1 for p in region_point.values() if p in altitude)}행")
    print(f"{DISTRICT_OUT.relative_to(ROOT)}  {sum(by_source.values())}행 {by_source}")
    if missing_regions:
        # 조용히 빠지면 그 구역만 고도 필터가 꺼진다 — 반드시 드러낸다.
        print(f"[확인 필요] 엑셀에 좌표 없는 구역 {len(missing_regions)}개: {missing_regions}")


if __name__ == "__main__":
    main()
